import pandas as pd, numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

T = {  # sheet name -> (pickle, description)
    "Activitate_lunara": (
        "uw",
        "Activitatea de asigurare pe lună × țară × linie de business × canal",
    ),
    "Daune": ("claims", "Eșantion de 25.000 de daune individuale"),
    "Buget": (
        "budget",
        "Bugetul de prime și combined ratio țintă, pe lună × regiune × linie",
    ),
    "Rezultate_grup": (
        "fin",
        "Contul de profit, capitaluri și indicatori bursieri ai grupului, trimestrial",
    ),
    "Investitii": (
        "invest",
        "Portofoliul de investiții pe clase de active, trimestrial",
    ),
    "Angajati": ("emp", "Toți angajații care au lucrat în companie între 2021 și 2025"),
    "Recrutare": ("rec", "Pâlnia de recrutare pe trimestru × departament × regiune"),
    "Sondaje_angajati": (
        "eng",
        "Sondajul anual de engagement pe departament × regiune",
    ),
    "Tinte_strategice": ("kpi", "Cei 10 KPI strategici: țintă vs. realizat, anual"),
    "Cursuri_valutare": ("fx", "Cursuri lunare: unități de monedă locală pentru 1 EUR"),
    "Tari": ("countries", "Țările în care activează grupul"),
}
data = {s: pd.read_pickle(p + ".pkl") for s, (p, _) in T.items()}
data["Tari"] = data["Tari"].drop(columns=["base_gwp_eur_m"])
data["Activitate_lunara"]["cat_event"] = data["Activitate_lunara"]["cat_event"].replace(
    "", None
)
data["Daune"]["cat_event"] = data["Daune"]["cat_event"].replace("", None)

# ---------------- data dictionary ----------------
D = []


def add(sheet, rows):
    for r in rows:
        D.append((sheet,) + r)


common = {
    "month": ("text", "AAAA-LL", "Luna la care se referă rândul"),
    "year": ("număr", "an", "Anul"),
    "quarter": (
        "text",
        "AAAA-Tn",
        "Trimestrul (Q1 = ian–mar, Q2 = apr–iun, Q3 = iul–sep, Q4 = oct–dec)",
    ),
    "region": (
        "text",
        "–",
        "Regiunea: Europe, North America, Asia-Pacific, Latin America",
    ),
    "country_code": (
        "text",
        "cod ISO",
        "Codul țării din 2 litere (DE = Germania, RO = România etc.); legătura cu foaia Tari",
    ),
    "line_of_business": (
        "text",
        "–",
        "Linia de business: Motor (auto), Property (locuințe și bunuri), Life (viață), Health (sănătate), Commercial (riscuri pentru firme)",
    ),
    "channel": (
        "text",
        "–",
        "Canalul de vânzare: Agents (agenți), Brokers (brokeri), Direct Online (direct, pe internet), Bancassurance (prin bănci partenere)",
    ),
    "currency": ("text", "cod ISO", "Moneda locală a țării (EUR, USD, RON etc.)"),
    "department": ("text", "–", "Departamentul angajatului"),
}


def c(sheet, cols):
    rows = []
    for col in cols:
        if isinstance(col, str):
            rows.append((col,) + common[col])
        else:
            rows.append(col)
    add(sheet, rows)


c(
    "Activitate_lunara",
    [
        "month",
        "year",
        "quarter",
        "region",
        "country_code",
        "line_of_business",
        "channel",
        "currency",
        (
            "gwp_eur",
            "număr",
            "EUR",
            "Prime brute subscrise (GWP): valoarea polițelor vândute în lună, înainte de reasigurare",
        ),
        (
            "gwp_local_ccy",
            "număr",
            "monedă locală",
            "Aceleași prime brute, în moneda țării (arată efectul cursului valutar)",
        ),
        (
            "net_earned_premium_eur",
            "număr",
            "EUR",
            "Prime nete câștigate (NEP): partea din prime aferentă lunii, după cedarea către reasigurători. Baza de calcul pentru rate",
        ),
        (
            "claims_incurred_eur",
            "număr",
            "EUR",
            "Daune întâmplate: daune plătite plus modificarea rezervelor pentru daunele încă nesoluționate",
        ),
        (
            "acquisition_costs_eur",
            "număr",
            "EUR",
            "Costuri de achiziție: comisioane pentru agenți, brokeri și bănci, plus marketing",
        ),
        (
            "admin_expenses_eur",
            "număr",
            "EUR",
            "Cheltuieli administrative: salarii din back-office, IT, sedii",
        ),
        ("policies_in_force", "număr", "polițe", "Polițe active la sfârșitul lunii"),
        ("new_policies", "număr", "polițe", "Polițe noi vândute în lună"),
        (
            "lapsed_policies",
            "număr",
            "polițe",
            "Polițe pierdute în lună (nereînnoite sau anulate)",
        ),
        (
            "cat_event",
            "text",
            "–",
            "Numele evenimentului catastrofal care a afectat luna (furtună, uragan, inundații); gol dacă nu a fost cazul",
        ),
    ],
)
c(
    "Daune",
    [
        ("claim_id", "text", "–", "Codul unic al daunei"),
        (
            "reported_date",
            "dată",
            "AAAA-LL-ZZ",
            "Data la care clientul a anunțat dauna",
        ),
        "region",
        "country_code",
        "line_of_business",
        "channel",
        (
            "cause",
            "text",
            "–",
            "Cauza daunei (Collision = coliziune, Theft = furt, Water damage = inundație din conducte, Storm = furtună etc.)",
        ),
        ("cat_event", "text", "–", "Evenimentul catastrofal asociat, dacă există"),
        ("claim_amount_eur", "număr", "EUR", "Suma cerută sau estimată inițial"),
        (
            "paid_amount_eur",
            "număr",
            "EUR",
            "Suma plătită efectiv; 0 pentru daunele deschise sau respinse",
        ),
        (
            "status",
            "text",
            "–",
            "Closed (închisă, plătită), Open (deschisă la 31.12.2025), Rejected (respinsă)",
        ),
        (
            "days_to_settle",
            "număr",
            "zile",
            "Zile de la anunțare până la soluționare; gol pentru daunele deschise",
        ),
        (
            "fraud_suspected",
            "adevărat/fals",
            "–",
            "TRUE dacă dauna a fost marcată ca suspectă de fraudă",
        ),
        (
            "customer_satisfaction_1_5",
            "număr",
            "scor 1–5",
            "Satisfacția clientului după soluționare (5 = foarte mulțumit); gol pentru daunele deschise",
        ),
    ],
)
c(
    "Buget",
    [
        "month",
        "year",
        "region",
        "line_of_business",
        ("budget_gwp_eur", "număr", "EUR", "Primele brute bugetate pentru lună"),
        (
            "target_combined_ratio",
            "număr",
            "fracție",
            "Combined ratio țintă (0,91 = 91%)",
        ),
    ],
)
c(
    "Rezultate_grup",
    [
        "quarter",
        ("gwp_eur", "număr", "EUR", "Prime brute subscrise în trimestru"),
        ("net_earned_premium_eur", "număr", "EUR", "Prime nete câștigate"),
        ("claims_incurred_eur", "număr", "EUR", "Daune întâmplate"),
        ("acquisition_costs_eur", "număr", "EUR", "Costuri de achiziție"),
        ("admin_expenses_eur", "număr", "EUR", "Cheltuieli administrative"),
        (
            "investment_income_eur",
            "număr",
            "EUR",
            "Venituri din investiții (dobânzi, dividende, chirii)",
        ),
        (
            "underwriting_result_eur",
            "număr",
            "EUR",
            "Rezultat tehnic = NEP − daune − costuri de achiziție − cheltuieli administrative",
        ),
        (
            "other_expenses_eur",
            "număr",
            "EUR",
            "Alte cheltuieli ale grupului (valoare negativă)",
        ),
        ("profit_before_tax_eur", "număr", "EUR", "Profit înainte de impozit"),
        ("tax_eur", "număr", "EUR", "Impozit pe profit (25%)"),
        ("net_income_eur", "număr", "EUR", "Profit net"),
        (
            "shares_outstanding",
            "număr",
            "acțiuni",
            "Numărul de acțiuni în circulație (scade în 2024–2025 prin răscumpărări)",
        ),
        ("eps_eur", "număr", "EUR/acțiune", "Profit pe acțiune (EPS) în trimestru"),
        (
            "dividends_paid_eur",
            "număr",
            "EUR",
            "Dividende plătite (în Q2, din profitul anului anterior)",
        ),
        (
            "shareholders_equity_eur",
            "număr",
            "EUR",
            "Capitaluri proprii la sfârșitul trimestrului",
        ),
        ("dps_eur", "număr", "EUR/acțiune", "Dividend pe acțiune (DPS)"),
        (
            "roe_annualised",
            "număr",
            "fracție",
            "Rentabilitatea capitalurilor proprii (ROE), anualizată: profitul trimestrului × 4 / capitaluri",
        ),
        (
            "solvency_ii_ratio",
            "număr",
            "fracție",
            "Rata de solvabilitate Solvency II (2,10 = 210%); minimum legal 100%",
        ),
        (
            "share_price_eur_end",
            "număr",
            "EUR",
            "Prețul acțiunii la sfârșitul trimestrului",
        ),
        (
            "market_cap_eur",
            "număr",
            "EUR",
            "Capitalizarea bursieră = preț × acțiuni în circulație",
        ),
    ],
)
c(
    "Investitii",
    [
        "quarter",
        (
            "asset_class",
            "text",
            "–",
            "Clasa de active: obligațiuni de stat, obligațiuni corporative, acțiuni, imobiliare, numerar",
        ),
        (
            "market_value_eur",
            "număr",
            "EUR",
            "Valoarea de piață la sfârșitul trimestrului",
        ),
        ("investment_income_eur", "număr", "EUR", "Venitul încasat în trimestru"),
        (
            "unrealised_gain_loss_eur",
            "număr",
            "EUR",
            "Câștig sau pierdere nerealizată din modificarea prețurilor (activul nu a fost vândut)",
        ),
    ],
)
c(
    "Angajati",
    [
        ("employee_id", "text", "–", "Codul unic al angajatului"),
        ("status", "text", "–", "Active (activ la 31.12.2025) sau Left (a plecat)"),
        ("gender", "text", "–", "F, M sau X (nedeclarat/non-binar)"),
        ("birth_year", "număr", "an", "Anul nașterii"),
        "region",
        "country_code",
        "department",
        (
            "job_level",
            "text",
            "–",
            "Nivelul ierarhic: L1 (junior) până la L6 (director), EX (executiv)",
        ),
        (
            "is_people_manager",
            "adevărat/fals",
            "–",
            "TRUE dacă are oameni în subordine",
        ),
        (
            "work_model",
            "text",
            "–",
            "Hybrid (hibrid), Office (la birou), Remote (de acasă)",
        ),
        (
            "hire_date",
            "dată",
            "AAAA-LL-ZZ",
            "Data angajării (poate fi dinainte de 2021)",
        ),
        (
            "exit_date",
            "dată",
            "AAAA-LL-ZZ",
            "Data plecării; gol pentru angajații activi",
        ),
        (
            "exit_type",
            "text",
            "–",
            "Voluntary (demisie), Involuntary (concediere), Retirement (pensionare)",
        ),
        ("exit_reason", "text", "–", "Motivul declarat al plecării"),
        (
            "base_salary_eur",
            "număr",
            "EUR/an",
            "Salariul de bază anual brut, convertit în EUR, la data plecării sau la 31.12.2025",
        ),
        (
            "bonus_target_pct",
            "număr",
            "fracție",
            "Bonusul țintă ca procent din salariu (0,12 = 12%)",
        ),
        (
            "last_performance_rating_1_5",
            "număr",
            "scor 1–5",
            "Ultima evaluare de performanță (3 = conform așteptărilor)",
        ),
    ],
)
c(
    "Recrutare",
    [
        "quarter",
        "department",
        "region",
        ("applications", "număr", "persoane", "Candidaturi primite"),
        ("screened", "număr", "persoane", "Candidați trecuți de preselecție"),
        ("interviews", "număr", "interviuri", "Interviuri susținute"),
        ("offers_made", "număr", "oferte", "Oferte de angajare făcute"),
        ("offers_accepted", "număr", "oferte", "Oferte acceptate"),
        ("hires", "număr", "persoane", "Angajări efective (coincid cu foaia Angajati)"),
        (
            "avg_time_to_hire_days",
            "număr",
            "zile",
            "Timpul mediu de la deschiderea postului până la acceptarea ofertei",
        ),
        (
            "cost_per_hire_eur",
            "număr",
            "EUR",
            "Costul mediu per angajare (anunțuri, recrutori, agenții)",
        ),
    ],
)
c(
    "Sondaje_angajati",
    [
        ("survey_year", "număr", "an", "Anul sondajului"),
        "department",
        "region",
        ("headcount_invited", "număr", "persoane", "Angajați invitați să răspundă"),
        ("participation_rate", "număr", "fracție", "Rata de participare (0,78 = 78%)"),
        (
            "engagement_score_0_100",
            "număr",
            "scor 0–100",
            "Indicele de engagement (implicare)",
        ),
        (
            "enps",
            "număr",
            "puncte −100…+100",
            "eNPS: % promotori minus % detractori la întrebarea „Ai recomanda compania ca loc de muncă?”",
        ),
        (
            "manager_support_0_100",
            "număr",
            "scor 0–100",
            "Sprijinul perceput din partea managerului",
        ),
        (
            "workload_balance_0_100",
            "număr",
            "scor 0–100",
            "Echilibrul volumului de muncă (mic = suprasolicitare)",
        ),
        (
            "career_growth_0_100",
            "număr",
            "scor 0–100",
            "Perspective de carieră percepute",
        ),
        ("pay_fairness_0_100", "număr", "scor 0–100", "Echitatea salarială percepută"),
    ],
)
c(
    "Tinte_strategice",
    [
        "year",
        ("kpi", "text", "–", "Numele indicatorului strategic"),
        ("owner", "text", "–", "Departamentul responsabil"),
        (
            "unit",
            "text",
            "–",
            "Unitatea: ratio sau pct (fracție), points (puncte), days (zile), score (scor)",
        ),
        ("direction", "text", "–", "Dacă o valoare mai mare sau mai mică e de dorit"),
        ("target", "număr", "vezi unit", "Ținta anului"),
        (
            "actual",
            "număr",
            "vezi unit",
            "Valoarea realizată, calculată din celelalte foi",
        ),
    ],
)
c(
    "Cursuri_valutare",
    [
        "month",
        "currency",
        (
            "local_per_eur",
            "număr",
            "monedă locală",
            "Câte unități de monedă locală valorează 1 EUR (media lunii)",
        ),
    ],
)
c(
    "Tari",
    ["country_code", ("country", "text", "–", "Denumirea țării"), "region", "currency"],
)
dict_df = pd.DataFrame(D, columns=["Foaie", "Coloană", "Tip", "Unitate", "Descriere"])

# ---------------- glossary ----------------
G = [
    # (termen RO, termen EN, domeniu, definiție, formula/exemplu)
    (
        "Prime brute subscrise",
        "Gross Written Premium (GWP)",
        "Asigurări – business",
        "Valoarea totală a polițelor vândute într-o perioadă, înainte de reasigurare. Arată mărimea business-ului.",
        "Exemplu: o poliță auto de 600 € vândută în martie adaugă 600 € la GWP-ul lunii martie.",
    ),
    (
        "Prime nete câștigate",
        "Net Earned Premium (NEP)",
        "Asigurări – business",
        "Partea din prime care corespunde perioadei analizate, după ce se scade partea cedată reasigurătorilor. Este venitul „real” al perioadei.",
        "O poliță anuală de 1.200 € vândută pe 1 ianuarie „câștigă” 100 € în fiecare lună.",
    ),
    (
        "Reasigurare",
        "Reinsurance",
        "Asigurări – business",
        "Asigurarea asigurătorului: o parte din riscuri și din prime se cedează altei companii, care preia o parte din daunele mari.",
        "InaVale cedează circa 6% din prime reasigurătorilor.",
    ),
    (
        "Linie de business",
        "Line of business",
        "Asigurări – business",
        "Categoria de produse de asigurare: auto, locuințe, viață, sănătate, riscuri pentru firme.",
        "–",
    ),
    (
        "Canal de distribuție",
        "Distribution channel",
        "Asigurări – business",
        "Modul în care se vând polițele: agenți, brokeri, online direct sau prin bănci.",
        "–",
    ),
    (
        "Bancassurance",
        "Bancassurance",
        "Asigurări – business",
        "Vânzarea asigurărilor prin rețeaua unei bănci partenere, de exemplu la acordarea unui credit.",
        "–",
    ),
    (
        "Broker",
        "Broker",
        "Asigurări – business",
        "Intermediar independent care reprezintă clientul și caută oferta potrivită la mai mulți asigurători.",
        "–",
    ),
    (
        "Polițe active",
        "Policies in force",
        "Asigurări – business",
        "Numărul de polițe valabile la un moment dat.",
        "–",
    ),
    (
        "Rata de pierdere a polițelor",
        "Lapse rate",
        "Asigurări – business",
        "Procentul de polițe care nu se reînnoiesc sau sunt anulate.",
        "Polițe pierdute / polițe active",
    ),
    (
        "Rata de retenție",
        "Retention rate",
        "Asigurări – business",
        "Procentul de clienți păstrați; complementul ratei de pierdere.",
        "1 − rata de pierdere",
    ),
    (
        "Subscriere",
        "Underwriting",
        "Asigurări – business",
        "Procesul prin care asigurătorul evaluează riscul și stabilește prețul poliței.",
        "–",
    ),
    (
        "Daune întâmplate",
        "Claims incurred",
        "Daune",
        "Costul daunelor dintr-o perioadă: sumele plătite plus modificarea rezervelor pentru daunele încă nesoluționate.",
        "Plăți + (rezerve la final − rezerve la început)",
    ),
    (
        "Rata daunei",
        "Loss ratio",
        "Daune",
        "Cât din venitul din prime se duce pe daune.",
        "Daune întâmplate / NEP. Exemplu: 700 € daune la 1.000 € NEP = 70%",
    ),
    (
        "Frecvența daunelor",
        "Claims frequency",
        "Daune",
        "Cât de des apar daunele raportat la numărul de polițe.",
        "Număr de daune / polițe active",
    ),
    (
        "Severitatea daunelor",
        "Claims severity",
        "Daune",
        "Valoarea medie a unei daune.",
        "Total daune / număr de daune",
    ),
    (
        "Eveniment catastrofal",
        "Catastrophe (CAT) event",
        "Daune",
        "Eveniment natural sau accidental care produce foarte multe daune simultan: furtuni, uragane, inundații, cutremure.",
        "–",
    ),
    (
        "Rezerve de daune",
        "Claims reserves",
        "Daune",
        "Bani puși deoparte pentru daunele întâmplate, dar încă neplătite.",
        "–",
    ),
    (
        "Timp de soluționare",
        "Days to settle",
        "Daune",
        "Zilele de la anunțarea daunei până la plata sau închiderea ei. Influențează direct satisfacția clienților.",
        "–",
    ),
    (
        "Fraudă suspectată",
        "Suspected fraud",
        "Daune",
        "Daună marcată pentru verificare suplimentară deoarece prezintă semnale de fraudă.",
        "–",
    ),
    (
        "Rata cheltuielilor",
        "Expense ratio",
        "Financiar",
        "Cât costă funcționarea companiei raportat la prime.",
        "(Costuri de achiziție + cheltuieli administrative) / NEP",
    ),
    (
        "Combined ratio",
        "Combined ratio",
        "Financiar",
        "Cel mai important indicator din asigurări. Sub 100% înseamnă că activitatea de asigurare aduce profit; peste 100% înseamnă pierdere, acoperită eventual din investiții.",
        "Rata daunei + rata cheltuielilor. Exemplu: 68% + 25% = 93%",
    ),
    (
        "Costuri de achiziție",
        "Acquisition costs",
        "Financiar",
        "Costurile vânzării polițelor: comisioane pentru intermediari și marketing.",
        "–",
    ),
    (
        "Rezultat tehnic",
        "Underwriting result",
        "Financiar",
        "Profitul din activitatea de asigurare, fără investiții.",
        "NEP − daune − cheltuieli",
    ),
    (
        "Profit net",
        "Net income",
        "Financiar",
        "Profitul după toate cheltuielile și impozite.",
        "–",
    ),
    (
        "Buget vs. realizat",
        "Budget vs. actual",
        "Financiar",
        "Compararea rezultatelor efective cu cele planificate.",
        "–",
    ),
    (
        "Abatere",
        "Variance",
        "Financiar",
        "Diferența dintre realizat și buget (sau față de anul anterior), în valoare absolută sau procentuală.",
        "Realizat − buget",
    ),
    (
        "Curs valutar",
        "Exchange rate (FX)",
        "Financiar",
        "Prețul unei monede exprimat în altă monedă. Pentru un grup global, schimbă valoarea în EUR a veniturilor locale.",
        "–",
    ),
    (
        "Efect valutar",
        "FX effect",
        "Financiar",
        "Partea din creșterea sau scăderea raportată în EUR care vine doar din modificarea cursului, nu din business.",
        "–",
    ),
    (
        "Capitaluri proprii",
        "Shareholders' equity",
        "Investitori",
        "Valoarea care aparține acționarilor: active minus datorii.",
        "–",
    ),
    (
        "Rentabilitatea capitalurilor proprii",
        "Return on equity (ROE)",
        "Investitori",
        "Cât profit produce compania pentru fiecare euro al acționarilor.",
        "Profit net / capitaluri proprii",
    ),
    (
        "Profit pe acțiune",
        "Earnings per share (EPS)",
        "Investitori",
        "Profitul net împărțit la numărul de acțiuni.",
        "Profit net / acțiuni în circulație",
    ),
    (
        "Dividend pe acțiune",
        "Dividend per share (DPS)",
        "Investitori",
        "Suma plătită acționarilor pentru fiecare acțiune.",
        "–",
    ),
    (
        "Rata de distribuire",
        "Payout ratio",
        "Investitori",
        "Ce parte din profit se plătește ca dividende.",
        "Dividende / profit net",
    ),
    (
        "Capitalizare bursieră",
        "Market capitalisation",
        "Investitori",
        "Valoarea de piață a companiei.",
        "Preț acțiune × acțiuni în circulație",
    ),
    (
        "Răscumpărare de acțiuni",
        "Share buyback",
        "Investitori",
        "Compania își cumpără propriile acțiuni, reducând numărul lor și crescând EPS-ul.",
        "–",
    ),
    (
        "Rata de solvabilitate",
        "Solvency II ratio",
        "Risc și reglementare",
        "Cât capital are compania față de capitalul minim cerut de reglementarea europeană Solvency II. Sub 100% e interzis; asigurătorii mari țintesc 150–220%.",
        "Capital disponibil / capital de solvabilitate cerut (SCR)",
    ),
    (
        "Solvency II",
        "Solvency II",
        "Risc și reglementare",
        "Directiva europeană care stabilește cât capital trebuie să dețină asigurătorii în funcție de riscurile lor.",
        "–",
    ),
    (
        "IFRS 17",
        "IFRS 17",
        "Risc și reglementare",
        "Standardul internațional de contabilitate pentru contractele de asigurare, în vigoare din 2023. Datele din acest set folosesc o versiune simplificată.",
        "–",
    ),
    (
        "Câștig/pierdere nerealizată",
        "Unrealised gain/loss",
        "Investiții",
        "Modificarea valorii unui activ care nu a fost încă vândut. De exemplu, obligațiunile pierd valoare când dobânzile cresc (2022).",
        "–",
    ),
    (
        "Randament",
        "Yield",
        "Investiții",
        "Venitul anual al unei investiții raportat la valoarea ei.",
        "Venit anual / valoare de piață",
    ),
    (
        "Clasă de active",
        "Asset class",
        "Investiții",
        "Grup de investiții similare: obligațiuni, acțiuni, imobiliare, numerar.",
        "–",
    ),
    (
        "Număr de angajați",
        "Headcount",
        "Resurse umane",
        "Numărul de angajați la o anumită dată.",
        "–",
    ),
    (
        "Rata plecărilor",
        "Attrition / turnover rate",
        "Resurse umane",
        "Procentul de angajați care pleacă într-o perioadă.",
        "Plecări / număr mediu de angajați",
    ),
    (
        "Plecări voluntare",
        "Voluntary attrition",
        "Resurse umane",
        "Plecări la inițiativa angajatului (demisii). Semnal important despre climatul intern.",
        "–",
    ),
    (
        "Engagement",
        "Employee engagement",
        "Resurse umane",
        "Gradul de implicare și atașament al angajaților față de companie și munca lor.",
        "Scor 0–100 din sondaj",
    ),
    (
        "eNPS",
        "Employee Net Promoter Score",
        "Resurse umane",
        "Cât de mult și-ar recomanda angajații compania ca loc de muncă. Variază între −100 și +100.",
        "% promotori (9–10) − % detractori (0–6)",
    ),
    (
        "Timp de angajare",
        "Time to hire",
        "Resurse umane",
        "Zilele necesare pentru ocuparea unui post.",
        "–",
    ),
    (
        "Pâlnie de recrutare",
        "Recruitment funnel",
        "Resurse umane",
        "Etapele succesive ale recrutării: candidaturi → preselecție → interviuri → oferte → angajări. La fiecare pas rămân mai puțini candidați.",
        "–",
    ),
    (
        "Diferență salarială de gen",
        "Gender pay gap",
        "Resurse umane",
        "Diferența procentuală dintre salariile medii ale femeilor și bărbaților, ideal comparată la același nivel și rol.",
        "(Salariu mediu B − salariu mediu F) / salariu mediu B",
    ),
    (
        "NPS clienți",
        "Net Promoter Score",
        "Clienți",
        "Cât de mult și-ar recomanda clienții asigurătorul.",
        "% promotori − % detractori",
    ),
    (
        "KPI",
        "Key Performance Indicator",
        "Analiză de date",
        "Indicator-cheie de performanță: o cifră aleasă pentru a urmări dacă un obiectiv se atinge.",
        "–",
    ),
    (
        "Țintă",
        "Target",
        "Analiză de date",
        "Valoarea pe care compania vrea să o atingă pentru un KPI.",
        "–",
    ),
    (
        "Tabel de fapte",
        "Fact table",
        "Analiză de date",
        "Tabel cu evenimente măsurabile (vânzări, daune, angajări); are multe rânduri și coloane numerice.",
        "Exemplu: Daune",
    ),
    (
        "Tabel de dimensiuni",
        "Dimension table",
        "Analiză de date",
        "Tabel care descrie contextul faptelor: când, unde, ce, cine.",
        "Exemplu: Tari",
    ),
    (
        "Schemă stea",
        "Star schema",
        "Analiză de date",
        "Model de date cu tabele de fapte în centru, legate de dimensiuni comune.",
        "–",
    ),
    (
        "Granularitate",
        "Grain",
        "Analiză de date",
        "Ce reprezintă un rând dintr-un tabel. Decide ce se poate calcula din el.",
        "Activitate_lunara: lună × țară × linie × canal",
    ),
    (
        "Detaliere",
        "Drill-down",
        "Analiză de date",
        "Coborârea de la o cifră agregată la componentele ei, de exemplu de la grup la regiune, țară, linie.",
        "–",
    ),
    (
        "Cohortă",
        "Cohort",
        "Analiză de date",
        "Grup de elemente cu un punct de plecare comun, urmărit în timp, de exemplu angajații angajați în 2022.",
        "–",
    ),
    (
        "Analiză Pareto",
        "Pareto analysis",
        "Analiză de date",
        "Identificarea puținelor elemente care produc majoritatea efectului (regula 80/20).",
        "–",
    ),
    (
        "Rată anuală compusă de creștere",
        "CAGR",
        "Analiză de date",
        "Creșterea medie anuală pe mai mulți ani.",
        "(Valoare finală / valoare inițială)^(1/ani) − 1",
    ),
    (
        "Anualizare",
        "Annualisation",
        "Analiză de date",
        "Transformarea unei valori trimestriale sau lunare într-una anuală, pentru comparație.",
        "Valoare trimestrială × 4",
    ),
    (
        "Date sintetice",
        "Synthetic data",
        "Analiză de date",
        "Date generate artificial, realiste statistic, dar care nu descriu persoane sau companii reale.",
        "Tot acest set de date",
    ),
]
gloss_df = pd.DataFrame(
    G,
    columns=[
        "Termen",
        "Termen în engleză",
        "Domeniu",
        "Definiție",
        "Formulă / exemplu",
    ],
)

# ---------------- README ----------------
readme = [
    ("InaVale Assurance Group – set de date sintetic", None),
    ("", None),
    (
        "Ce este",
        "Date fictive, dar realiste, ale unui grup global de asigurări, create pentru un dashboard de analiză pe mai multe layere. Nicio persoană, companie sau cifră nu este reală.",
    ),
    ("Perioada", "Ianuarie 2021 – decembrie 2025 (5 ani)"),
    (
        "Acoperire",
        "11 țări în 4 regiuni, 5 linii de business, 4 canale de vânzare, 11 departamente",
    ),
    (
        "Moneda de raportare",
        "EUR. Sumele locale sunt convertite cu foaia Cursuri_valutare.",
    ),
    (
        "Simplificări",
        "Contabilitatea de asigurări (IFRS 17) și calculul solvabilității (Solvency II) sunt simplificate, dar păstrează logica reală a indicatorilor.",
    ),
    (
        "Legături între foi",
        "Foile se leagă prin coloanele comune: month / quarter / year (timp), region și country_code (geografie), line_of_business (produs), channel (canal), department (departament).",
    ),
    (
        "Valorile din celule",
        "Numele coloanelor și valorile categoriale sunt în engleză, ca în sistemele reale. Traducerea și explicația fiecăreia sunt în foaia Dictionar_date.",
    ),
    ("Glosar", "Foaia Glosar explică toți termenii tehnici, grupați pe domenii."),
    ("", None),
    ("Foaie", "Conținut și număr de rânduri"),
]
for s, (p, desc) in T.items():
    readme.append((s, f"{desc} – {len(data[s]):,} rânduri".replace(",", ".")))

import os

out = os.environ.get("XLSX_OUT", "../data/InaVale_Assurance_date.xlsx")
with pd.ExcelWriter(out, engine="openpyxl") as w:
    pd.DataFrame(readme).to_excel(w, sheet_name="Citeste_ma", index=False, header=False)
    dict_df.to_excel(w, sheet_name="Dictionar_date", index=False)
    gloss_df.to_excel(w, sheet_name="Glosar", index=False)
    for s in T:
        data[s].to_excel(w, sheet_name=s, index=False)

# ---------------- formatting ----------------
wb = load_workbook(out)
F = Font(name="Arial", size=10)
FB = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HF = PatternFill("solid", fgColor="1F3A5F")
money_cols = lambda n: n.endswith("_eur") or n.endswith("_ccy")
pct_cols = {
    "target_combined_ratio",
    "roe_annualised",
    "solvency_ii_ratio",
    "participation_rate",
    "bonus_target_pct",
}
for ws in wb.worksheets:
    if ws.title == "Citeste_ma":
        for row in ws.iter_rows():
            for cell in row:
                cell.font = F
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws["A1"].font = Font(name="Arial", size=14, bold=True)
        for r in range(2, ws.max_row + 1):
            ws.cell(r, 1).font = Font(name="Arial", size=10, bold=True)
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 110
        continue
    headers = [c.value for c in ws[1]]
    for c in ws[1]:
        c.font = FB
        c.fill = HF
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    text_sheet = ws.title in ("Dictionar_date", "Glosar")
    for j, h in enumerate(headers, 1):
        col = get_column_letter(j)
        if text_sheet:
            width = {
                "Descriere": 90,
                "Definiție": 80,
                "Formulă / exemplu": 55,
                "Termen": 30,
                "Termen în engleză": 28,
                "Coloană": 30,
                "Foaie": 20,
                "Domeniu": 22,
            }.get(h, 14)
        else:
            width = max(12, min(34, len(str(h)) + 3))
        ws.column_dimensions[col].width = width
        fmt = None
        if not text_sheet:
            if h in pct_cols:
                fmt = "0.0%"
            elif h in ("eps_eur", "dps_eur", "share_price_eur_end", "local_per_eur"):
                fmt = "#,##0.00"
            elif h in ("target", "actual"):
                fmt = "0.000"
            elif money_cols(h) or h in (
                "shares_outstanding",
                "policies_in_force",
                "new_policies",
                "lapsed_policies",
                "applications",
            ):
                fmt = "#,##0"
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(r, j)
            cell.font = F
            if fmt:
                cell.number_format = fmt
            if text_sheet:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        if h in ("reported_date", "hire_date", "exit_date"):
            for r in range(2, ws.max_row + 1):
                ws.cell(r, j).number_format = "yyyy-mm-dd"
    ws.row_dimensions[1].height = 30
wb.save(out)
print("ok", len(dict_df), "coloane documentate;", len(gloss_df), "termeni")
